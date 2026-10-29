"""Convert a We360 workforce workbook (.xlsx) into a dashboard client folder.

    python3 scripts/import_workbook.py <workbook.xlsx> data/<client-id>

Reads the workbook's saved (calculated) values, so it must have been saved by Excel or
recalculated. Produces employees.csv, daily_activity.csv, app_usage.csv and client_config.json
in the format described in data/README.md.
"""
import csv, json, os, sys
from openpyxl import load_workbook

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
wb = load_workbook(src, read_only=True, data_only=True)


def table(name):
    rows = wb[name].iter_rows(values_only=True)
    head = [h for h in next(rows)]
    for r in rows:
        if r and r[0] is not None:
            yield dict(zip(head, r))


def hms(hours):
    sec = int(round((hours or 0) * 3600))
    return f'{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}'


def write(fname, header, rows):
    with open(os.path.join(out, fname), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(header)
        n = 0
        for r in rows:
            w.writerow(r); n += 1
    print(f'{fname}: {n:,} rows')


# ---------------------------------------------------------------- employees
write('employees.csv',
      ['Employee ID', 'Name', 'Department', 'Team', 'Location', 'Joining Date', 'Excluded'],
      ([e['employee_id'], e['name'], e['team'], e['team'], e.get('location') or '',
        e['joining_date'].strftime('%Y-%m-%d') if e.get('joining_date') else '', 'No']
       for e in table('Employees')))

# ---------------------------------------------------------------- daily ledger
write('daily_activity.csv',
      ['Employee ID', 'Date', 'Present', 'Absence Reason', 'Online Time', 'Active Time', 'Productive Time',
       'Neutral Time', 'Unproductive Time', 'Punch Span'],
      ([d['employee_id'], d['date'].strftime('%Y-%m-%d'), 'Yes' if d['status'] == 'Present' else 'No',
        d.get('absence_reason') or '', hms(d['online_h']), hms(d['active_h']), hms(d['productive_h']),
        hms(d['neutral_h']), hms(d['unproductive_h']), hms(d['working_h']) if d['status'] == 'Present' else '']
       for d in table('Daily')))

# ---------------------------------------------------------------- app usage (monthly buckets)
write('app_usage.csv',
      ['Employee ID', 'Date', 'App/URL', 'Type', 'Productivity', 'Category', 'Active Time'],
      ([u['employee_id'], f"{u['month']}-01", u['item'], 'URL' if u['kind'] == 'url' else 'App',
        u['classification'].lower(), u['category'], hms(u['active_h'])]
       for u in table('App_Usage')))

# ---------------------------------------------------------------- client config
settings = {s['setting']: s['value'] for s in table('Settings')}
cc = {c['key']: c['value'] for c in table('Client_Config')}
split = lambda v: [x.strip() for x in str(v).split(',') if x.strip()]
ai_tools = {}
for a in table('Apps_Catalog'):
    if a.get('is_ai') and a.get('ai_tool'):
        ai_tools.setdefault(a['ai_tool'], []).append(str(a['item']).lower())

bench = float(settings['capacity_benchmark_h'])
cmin, cmax = float(settings['capacity_min_h']), float(settings['capacity_max_h'])
config = {
    'clientName': cc.get('client_name', 'Client'),
    'vendorName': 'We360.ai',
    'orgLevel': 'team',   # teams only - no departments in this workbook
    'excludeFromTopApps': split(cc.get('exclude_from_top_apps', '')),
    'nonWorkCategories': split(cc.get('non_work_categories', '')),
    'aiTools': ai_tools,
    'aiHubTool': cc.get('ai_hub_tool'),
    'minPresentDaysForLists': 15,
    'minGroupSize': 3,
    'defaults': {
        'capacity': {
            'metric': settings.get('capacity_basis', 'online'),
            'benchmark': bench,
            'variancePct': round((cmax - cmin) / 2 / bench * 100, 1),
            'min': cmin,
            'max': cmax,
        },
        'insightRules': {
            'severeCategories': split(cc.get('severe_non_work_categories', 'Gaming, Adult, Gambling')),
        },
    },
}
with open(os.path.join(out, 'client_config.json'), 'w', encoding='utf-8') as f:
    json.dump(config, f, indent=2, ensure_ascii=False)
    f.write('\n')
print('client_config.json written:', config['clientName'], '| capacity', cmin, '-', cmax, '| AI tools', len(ai_tools))
